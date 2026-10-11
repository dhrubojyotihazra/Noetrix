"""Unit tests for /api/parse endpoint hardening (Item 15 & Item 19)."""

import io
import pytest
from fastapi.testclient import TestClient
from docx import Document
from reportlab.pdfgen import canvas

from app.main import app


@pytest.fixture
def client():
    return TestClient(app)


def test_parse_pdf_success(client):
    """PDF upload succeeds and returns ParsedResume."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "Jane Doe")
    c.drawString(100, 730, "jane@example.com | 555-123-4567")
    c.drawString(100, 700, "EXPERIENCE")
    c.drawString(100, 680, "Engineer at TechCorp 2022-2024")
    c.save()
    buf.seek(0)

    resp = client.post(
        "/api/parse",
        files={"file": ("resume.pdf", buf.getvalue(), "application/pdf")}
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["filename"] == "resume.pdf"
    assert "jane@example.com" in data["contact"]["email"]


def test_parse_docx_success(client):
    """DOCX upload succeeds and returns ParsedResume."""
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("John Smith\njohn.smith@example.com | 555-987-6543")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Developer at Startup 2021-2023")
    doc.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/parse",
        files={"file": ("resume.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["filename"] == "resume.docx"
    assert "john.smith@example.com" in data["contact"]["email"]


def test_parse_txt_success(client):
    """TXT upload succeeds and returns ParsedResume."""
    text_content = b"Alice Lee\nalice@example.com | 555-333-4444\nEXPERIENCE\nData Scientist 2020-2024"
    resp = client.post(
        "/api/parse",
        files={"file": ("resume.txt", text_content, "text/plain")}
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["filename"] == "resume.txt"
    assert "alice@example.com" in data["contact"]["email"]


def test_parse_bad_extension_rejected(client):
    """Unsupported extensions like .png or .jpg return 400."""
    resp = client.post(
        "/api/parse",
        files={"file": ("resume.png", b"fake image bytes", "image/png")}
    )
    assert resp.status_code == 400
    assert "Unsupported file format" in resp.json()["detail"]


def test_parse_legacy_doc_rejected_with_clear_error(client):
    """Legacy .doc files must be rejected with a clear informative error message."""
    resp = client.post(
        "/api/parse",
        files={"file": ("resume.doc", b"fake doc bytes", "application/msword")}
    )
    assert resp.status_code == 400
    assert ".doc" in resp.json()["detail"]


def test_parse_oversized_file_rejected(client, monkeypatch):
    """Files exceeding MAX_UPLOAD_SIZE_MB return 413 Payload Too Large."""
    # Set limit to 0.001 MB (~1 KB)
    monkeypatch.setenv("MAX_UPLOAD_SIZE_MB", "0.001")
    large_payload = b"A" * 5000  # 5 KB > 1 KB

    resp = client.post(
        "/api/parse",
        files={"file": ("resume.txt", large_payload, "text/plain")}
    )
    assert resp.status_code == 413
    assert "maximum upload size" in resp.json()["detail"].lower()
