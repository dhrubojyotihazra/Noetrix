"""Extract structured text and formatting metrics from DOCX files using python-docx."""

import io
import docx

from app.schemas.resume import ParsedResume, WorkEntry, EduEntry
from app.services.parser.contact_extractor import extract_contact_info
from app.services.parser.section_detector import split_into_sections
from app.nlp.tokenizer import word_boundary_match
from app.nlp.tfidf import ALL_TAXONOMY_SKILLS


def parse_docx_bytes(file_bytes: bytes, filename: str = "resume.docx") -> ParsedResume:
    """Parse resume DOCX into structured ParsedResume object."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs_text = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            paragraphs_text.append(t)

    # Check tables
    has_tables = len(doc.tables) > 0
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs_text.append(row_text)

    raw_text = "\n".join(paragraphs_text).strip()
    contact = extract_contact_info(raw_text)
    sections = split_into_sections(raw_text)

    # Work experience
    work_entries: list[WorkEntry] = []
    exp_lines = sections.get("experience", [])
    if exp_lines:
        current_work = WorkEntry(id="exp_1", title=exp_lines[0])
        for line in exp_lines[1:]:
            if line.startswith(("-", "•", "*", "–")) or len(line) > 30:
                current_work.bullets.append(line.lstrip("-•*– ").strip())
            else:
                if current_work.bullets:
                    work_entries.append(current_work)
                    current_work = WorkEntry(id=f"exp_{len(work_entries)+1}", title=line)
                else:
                    current_work.company = line
        if current_work.bullets or current_work.title:
            work_entries.append(current_work)

    # Education
    edu_entries: list[EduEntry] = []
    edu_lines = sections.get("education", [])
    if edu_lines:
        edu_entries.append(EduEntry(
            id="edu_1",
            institution=edu_lines[0],
            degree=edu_lines[1] if len(edu_lines) > 1 else "",
            bullets=edu_lines[2:]
        ))

    # Skills
    skills_lines = sections.get("skills", [])
    extracted_skills: set[str] = set()
    skills_text = " ".join(skills_lines).lower()
    for skill in ALL_TAXONOMY_SKILLS:
        if word_boundary_match(skill, skills_text):
            extracted_skills.add(skill)

    if len(extracted_skills) < 5:
        lower_raw = raw_text.lower()
        for skill in ALL_TAXONOMY_SKILLS:
            if word_boundary_match(skill, lower_raw):
                extracted_skills.add(skill)

    sections_detected = [k for k, v in sections.items() if len(v) > 0 and k != "header" and k != "other"]

    return ParsedResume(
        filename=filename,
        raw_text=raw_text,
        contact=contact,
        summary="\n".join(sections.get("summary", [])),
        sections_detected=sections_detected,
        work_experience=work_entries,
        education=edu_entries,
        skills=sorted(list(extracted_skills)),
        projects=[],
        certifications=sections.get("certifications", []),
        parse_confidence=1.0 if len(sections_detected) >= 3 else 0.7,
        detected_columns=1,
        has_tables=has_tables
    )
